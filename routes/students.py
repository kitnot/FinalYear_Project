from flask import (
    Blueprint,
    render_template,
    request,
    redirect,
    url_for,
    flash,
    jsonify
)

from flask_login import login_required

from extensions import db

from models.student import Student
from models.face_embedding import FaceEmbedding
from models.attendance import Attendance

from face_engine.detector import FaceDetector
from face_engine.recognizer import FaceRecognizer

import cv2
import numpy as np
import csv
import io

from openpyxl import load_workbook


students = Blueprint(
    "students",
    __name__,
    url_prefix="/students"
)


detector = FaceDetector()
recognizer = FaceRecognizer()


# ============================================================
# FACE FEATURE EXTRACTION
# ============================================================

def extract_feature(image_bytes):

    array = np.frombuffer(
        image_bytes,
        dtype=np.uint8
    )

    frame = cv2.imdecode(
        array,
        cv2.IMREAD_COLOR
    )

    if frame is None:
        return None, "Invalid image."

    faces = detector.detect(frame)

    if len(faces) == 0:
        return None, "No face detected."

    if len(faces) > 1:
        return None, (
            "Multiple faces detected. "
            "Keep only one person in the camera."
        )

    feature = recognizer.get_feature(
        frame,
        faces[0]
    )

    return feature, None


# ============================================================
# STUDENTS LIST
# ============================================================

@students.route("/")
@login_required
def index():

    all_students = (
        Student.query
        .order_by(Student.student_id.asc())
        .all()
    )

    return render_template(
        "students.html",
        students=all_students
    )


# ============================================================
# ADD SINGLE STUDENT
# ============================================================

@students.route("/add", methods=["GET", "POST"])
@login_required
def add_student():

    if request.method == "POST":

        student_id = request.form.get(
            "student_id",
            ""
        ).strip()

        name = request.form.get(
            "name",
            ""
        ).strip()

        if not student_id or not name:

            flash(
                "Student ID and Name are required."
            )

            return redirect(
                url_for("students.add_student")
            )

        existing = Student.query.filter_by(
            student_id=student_id
        ).first()

        if existing:

            flash(
                "Student ID already exists."
            )

            return redirect(
                url_for("students.add_student")
            )

        student = Student(

            student_id=student_id,

            name=name,

            email=request.form.get(
                "email",
                ""
            ).strip(),

            phone=request.form.get(
                "phone",
                ""
            ).strip(),

            course=request.form.get(
                "course",
                ""
            ).strip(),

            semester=request.form.get(
                "semester",
                ""
            ).strip(),

            division=request.form.get(
                "division",
                ""
            ).strip()
        )

        db.session.add(student)

        db.session.commit()

        flash(
            f"Student '{name}' added successfully."
        )

        return redirect(
            url_for("students.index")
        )

    return render_template(
        "add_student.html"
    )


# ============================================================
# EDIT STUDENT
# ============================================================

@students.route(
    "/edit/<int:student_id>",
    methods=["GET", "POST"]
)
@login_required
def edit_student(student_id):

    student = Student.query.get_or_404(
        student_id
    )

    if request.method == "POST":

        new_id = request.form.get(
            "student_id",
            ""
        ).strip()

        name = request.form.get(
            "name",
            ""
        ).strip()

        if not new_id or not name:

            flash(
                "Student ID and Name are required."
            )

            return redirect(
                url_for(
                    "students.edit_student",
                    student_id=student.id
                )
            )

        duplicate = Student.query.filter(
            Student.student_id == new_id,
            Student.id != student.id
        ).first()

        if duplicate:

            flash(
                "Another student already uses that Student ID."
            )

            return redirect(
                url_for(
                    "students.edit_student",
                    student_id=student.id
                )
            )

        student.student_id = new_id

        student.name = name

        student.email = request.form.get(
            "email",
            ""
        ).strip()

        student.phone = request.form.get(
            "phone",
            ""
        ).strip()

        student.course = request.form.get(
            "course",
            ""
        ).strip()

        student.semester = request.form.get(
            "semester",
            ""
        ).strip()

        student.division = request.form.get(
            "division",
            ""
        ).strip()

        db.session.commit()

        flash(
            "Student updated successfully."
        )

        return redirect(
            url_for("students.index")
        )

    return render_template(
        "edit_student.html",
        student=student
    )


# ============================================================
# DELETE STUDENT
# ============================================================

@students.route(
    "/delete/<int:student_id>",
    methods=["POST"]
)
@login_required
def delete_student(student_id):

    student = Student.query.get_or_404(
        student_id
    )

    FaceEmbedding.query.filter_by(
        student_id=student.id
    ).delete(
        synchronize_session=False
    )

    Attendance.query.filter_by(
        student_id=student.id
    ).delete(
        synchronize_session=False
    )

    db.session.delete(student)

    db.session.commit()

    flash(
        f"Student '{student.name}' deleted."
    )

    return redirect(
        url_for("students.index")
    )


# ============================================================
# INDIVIDUAL FACE REGISTRATION PAGE
# ============================================================

@students.route(
    "/<int:student_id>/register-face"
)
@login_required
def register_face(student_id):

    student = Student.query.get_or_404(
        student_id
    )

    count = FaceEmbedding.query.filter_by(
        student_id=student.id
    ).count()

    return render_template(
        "register_face.html",
        student=student,
        sample_count=count
    )


# ============================================================
# CAPTURE FACE
# ============================================================

@students.route(
    "/<int:student_id>/capture-face",
    methods=["POST"]
)
@login_required
def capture_face(student_id):

    student = Student.query.get_or_404(
        student_id
    )

    image = request.files.get(
        "image"
    )

    if not image:

        return jsonify(
            success=False,
            message="No camera image received."
        ), 400

    feature, error = extract_feature(
        image.read()
    )

    if error:

        return jsonify(
            success=False,
            message=error
        ), 400

    embedding = FaceEmbedding(
        student_id=student.id,
        embedding=feature.astype(
            np.float32
        ).tobytes()
    )

    db.session.add(
        embedding
    )

    db.session.commit()

    count = FaceEmbedding.query.filter_by(
        student_id=student.id
    ).count()

    return jsonify(
        success=True,
        message=f"Face sample {count} saved.",
        sample_count=count
    )


# ============================================================
# IMPORT WHOLE CLASS
# ============================================================

@students.route(
    "/import-class",
    methods=["GET", "POST"]
)
@login_required
def import_class():

    if request.method == "GET":

        return render_template(
            "import_class.html"
        )

    uploaded = request.files.get(
        "class_file"
    )

    if not uploaded or not uploaded.filename:

        flash(
            "Please select a CSV or XLSX class file."
        )

        return redirect(
            url_for("students.import_class")
        )

    filename = uploaded.filename.lower()

    if not filename.endswith(
        (".csv", ".xlsx")
    ):

        flash(
            "Only CSV and XLSX files are supported."
        )

        return redirect(
            url_for("students.import_class")
        )

    try:

        rows = []

        # ====================================================
        # CSV
        # ====================================================

        if filename.endswith(".csv"):

            text = uploaded.read().decode(
                "utf-8-sig"
            )

            rows = list(
                csv.DictReader(
                    io.StringIO(text)
                )
            )

        # ====================================================
        # EXCEL
        # ====================================================

        else:

            workbook = load_workbook(
                uploaded,
                read_only=True,
                data_only=True
            )

            sheet = workbook.active

            values = list(
                sheet.iter_rows(
                    values_only=True
                )
            )

            if not values:

                flash(
                    "The Excel file is empty."
                )

                return redirect(
                    url_for(
                        "students.import_class"
                    )
                )

            headers = []

            for value in values[0]:

                if value is None:

                    headers.append("")

                else:

                    headers.append(
                        str(value)
                        .strip()
                        .lower()
                        .replace(" ", "_")
                    )

            for values_row in values[1:]:

                row = {}

                for index, header in enumerate(
                    headers
                ):

                    if not header:
                        continue

                    if index >= len(
                        values_row
                    ):

                        row[header] = ""

                    else:

                        value = values_row[index]

                        if value is None:
                            value = ""

                        row[header] = str(
                            value
                        ).strip()

                rows.append(row)

        # ====================================================
        # IMPORT
        # ====================================================

        added = 0
        skipped = 0
        errors = []

        for row_number, row in enumerate(
            rows,
            start=2
        ):

            cleaned = {}

            for key, value in row.items():

                key = (
                    str(key)
                    .strip()
                    .lower()
                    .replace(" ", "_")
                )

                if value is None:
                    value = ""

                cleaned[key] = str(
                    value
                ).strip()

            student_id = cleaned.get(
                "student_id",
                ""
            )

            name = cleaned.get(
                "name",
                ""
            )

            if not student_id:

                errors.append(
                    f"Row {row_number}: missing Student ID"
                )

                continue

            if not name:

                errors.append(
                    f"Row {row_number}: missing Name"
                )

                continue

            existing = Student.query.filter_by(
                student_id=student_id
            ).first()

            if existing:

                skipped += 1

                continue

            student = Student(

                student_id=student_id,

                name=name,

                email=cleaned.get(
                    "email",
                    ""
                ),

                phone=cleaned.get(
                    "phone",
                    ""
                ),

                course=cleaned.get(
                    "course",
                    ""
                ),

                semester=cleaned.get(
                    "semester",
                    ""
                ),

                division=cleaned.get(
                    "division",
                    ""
                )
            )

            db.session.add(
                student
            )

            added += 1

        db.session.commit()

        flash(
            f"Import complete: "
            f"{added} added, "
            f"{skipped} already existed, "
            f"{len(errors)} errors."
        )

        for error in errors[:10]:

            flash(error)

        return redirect(
            url_for("students.index")
        )

    except Exception as exc:

        db.session.rollback()

        flash(
            f"Import failed: {exc}"
        )

        return redirect(
            url_for("students.import_class")
        )


# ============================================================
# BATCH FACE REGISTRATION
# ============================================================

@students.route(
    "/batch-face-registration"
)
@login_required
def batch_face_registration():

    data = []

    all_students = (
        Student.query
        .order_by(Student.student_id.asc())
        .all()
    )

    for student in all_students:

        count = FaceEmbedding.query.filter_by(
            student_id=student.id
        ).count()

        data.append({

            "id": student.id,

            "student_id": student.student_id,

            "name": student.name,

            "sample_count": count
        })

    return render_template(
        "batch_face_registration.html",
        students=data
    )